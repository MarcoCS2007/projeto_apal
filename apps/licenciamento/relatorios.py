from collections import Counter, defaultdict
from datetime import timedelta
from decimal import Decimal
from io import BytesIO

from django.db.models import Count, Q
from django.utils import timezone

from apps.espacos.models import Endereco, PontoOcupacao, StatusOcupacao
from apps.fiscalizacao.models import OcorrenciaInspecao
from apps.usuarios.models import Ambulante, Genero

from .models import STATUS_FILA, LicencaAlvara, StatusLicenca
from .services import marcar_licencas_vencidas

PERIODO_INFRACOES_DIAS = 30
LIMIAR_LOTACAO = 90
CIDADE_SEDE = "Vitória da Conquista"

FAIXAS_ETARIAS = (
    ("18-24", 18, 24, "18–24 anos"),
    ("25-45", 25, 45, "25–45 anos"),
    ("46-60", 46, 60, "46–60 anos"),
    ("60+", 61, 200, "60 anos ou mais"),
)

FAIXAS_RENDA = (
    ("ate_1500", Decimal("0"), Decimal("1500"), "Até R$ 1.500"),
    ("1501_3000", Decimal("1500.01"), Decimal("3000"), "R$ 1.501 a R$ 3.000"),
    ("3001_5000", Decimal("3000.01"), Decimal("5000"), "R$ 3.001 a R$ 5.000"),
    ("acima_5000", Decimal("5000.01"), None, "Acima de R$ 5.000"),
)

GRUPO_ESCOLARIDADE = {
    "fundamental_incompleto": "Ensino Fundamental",
    "fundamental_completo": "Ensino Fundamental",
    "medio_incompleto": "Ensino Médio",
    "medio_completo": "Ensino Médio",
    "superior": "Ensino Superior",
    "analfabeto": "Sem instrução formal",
}


def _aplicar_origem(qs, origem, campo_ambulante="ambulante"):
    prefixo = f"{campo_ambulante}__" if campo_ambulante else ""
    if origem == "residentes":
        return qs.filter(**{f"{prefixo}tipo_atuacao": "fixo"})
    if origem == "itinerantes":
        return qs.filter(**{f"{prefixo}tipo_atuacao": "movel"})
    return qs


def _percentual(parte, total):
    if not total:
        return 0.0
    return round((parte / total) * 100, 1)


def _idade(nascimento, hoje=None):
    if not nascimento:
        return None
    hoje = hoje or timezone.localdate()
    anos = hoje.year - nascimento.year
    if (hoje.month, hoje.day) < (nascimento.month, nascimento.day):
        anos -= 1
    return anos


def _faixa_etaria(idade):
    if idade is None:
        return "", "Não informado"
    for chave, minimo, maximo, rotulo in FAIXAS_ETARIAS:
        if minimo <= idade <= maximo:
            return chave, rotulo
    if idade < 18:
        return "ate_17", "Até 17 anos"
    return "60+", "60 anos ou mais"


def _grupo_escolaridade(valor):
    if not valor:
        return "Não informado"
    if valor in GRUPO_ESCOLARIDADE:
        return GRUPO_ESCOLARIDADE[valor]
    texto = valor.lower()
    if "fundamental" in texto:
        return "Ensino Fundamental"
    if "médio" in texto or "medio" in texto:
        return "Ensino Médio"
    if "superior" in texto:
        return "Ensino Superior"
    if "analfabet" in texto:
        return "Sem instrução formal"
    return valor


def _faixa_renda(valor):
    if valor is None:
        return "", "Não informado"
    renda = Decimal(valor)
    for chave, minimo, maximo, rotulo in FAIXAS_RENDA:
        if maximo is None and renda >= minimo:
            return chave, rotulo
        if maximo is not None and minimo <= renda <= maximo:
            return chave, rotulo
    return "", "Não informado"


def _contar(itens):
    total = sum(item["total"] for item in itens)
    for item in itens:
        item["percentual"] = _percentual(item["total"], total)
    return itens, total


def filtros_do_request(request):
    return {
        "bairro": request.GET.get("bairro", ""),
        "origem": request.GET.get("origem", ""),
        "periodo_dias": request.GET.get("periodo", PERIODO_INFRACOES_DIAS),
        "cidade": request.GET.get("cidade", ""),
        "genero": request.GET.get("genero", ""),
        "faixa_etaria": request.GET.get("faixa", ""),
        "escolaridade": request.GET.get("escolaridade", ""),
    }


def indicadores_gerenciais(
    bairro="",
    origem="",
    periodo_dias=PERIODO_INFRACOES_DIAS,
    cidade="",
    genero="",
    faixa_etaria="",
    escolaridade="",
):
    """Números reais de licença, ocupação, infração, categoria e retrato social."""
    marcar_licencas_vencidas()
    bairro = (bairro or "").strip()
    origem = (origem or "").strip()
    cidade = (cidade or "").strip()
    genero = (genero or "").strip()
    faixa_etaria = (faixa_etaria or "").strip()
    escolaridade = (escolaridade or "").strip()
    try:
        periodo_dias = max(1, int(periodo_dias or PERIODO_INFRACOES_DIAS))
    except (TypeError, ValueError):
        periodo_dias = PERIODO_INFRACOES_DIAS

    ambulantes = _filtrar_ambulantes(
        bairro=bairro,
        origem=origem,
        cidade=cidade,
        genero=genero,
        faixa_etaria=faixa_etaria,
        escolaridade=escolaridade,
    )
    ids_ambulantes = list(ambulantes.values_list("pk", flat=True))

    licencas = LicencaAlvara.objects.all()
    if bairro:
        licencas = licencas.filter(ponto_ocupacao__bairro__iexact=bairro)
    licencas = _aplicar_origem(licencas, origem)
    if cidade or genero or faixa_etaria or escolaridade:
        licencas = licencas.filter(ambulante_id__in=ids_ambulantes or [0])

    hoje = timezone.localdate()
    ativas = licencas.filter(status=StatusLicenca.ATIVO)
    vencidas = licencas.filter(status=StatusLicenca.VENCIDO)
    pendentes = licencas.filter(status__in=STATUS_FILA)
    suspensas = licencas.filter(status=StatusLicenca.SUSPENSO)
    vencendo = ativas.filter(
        data_vencimento__gte=hoje,
        data_vencimento__lte=hoje + timedelta(days=30),
    )

    pontos = PontoOcupacao.objects.filter(ativo=True)
    if bairro:
        pontos = pontos.filter(bairro__iexact=bairro)
    ocupacao_qs = (
        pontos.values("bairro")
        .annotate(
            total=Count("id"),
            ocupados=Count("id", filter=Q(status_ocupacao=StatusOcupacao.OCUPADO)),
            livres=Count("id", filter=Q(status_ocupacao=StatusOcupacao.LIVRE)),
        )
        .order_by("bairro")
    )
    ativas_por_bairro = {
        row["ponto_ocupacao__bairro"]: row["total"]
        for row in ativas.filter(ponto_ocupacao__isnull=False)
        .values("ponto_ocupacao__bairro")
        .annotate(total=Count("id"))
        if row["ponto_ocupacao__bairro"]
    }
    ocupacao = []
    alerta_lotacao = None
    for row in ocupacao_qs:
        nome = row["bairro"] or "Não informado"
        total = row["total"]
        ocupados = row["ocupados"]
        percentual = _percentual(ocupados, total)
        item = {
            "bairro": nome,
            "total": total,
            "ocupados": ocupados,
            "livres": row["livres"],
            "percentual": percentual,
            "licencas_ativas": ativas_por_bairro.get(row["bairro"], 0),
        }
        ocupacao.append(item)
        if percentual >= LIMIAR_LOTACAO and alerta_lotacao is None:
            alerta_lotacao = item

    categorias = [
        {
            "nome": row["categoria_produto__nome_categoria"] or "Sem categoria",
            "total": row["total"],
        }
        for row in ativas.values("categoria_produto__nome_categoria")
        .annotate(total=Count("id"))
        .order_by("-total", "categoria_produto__nome_categoria")
    ]

    desde = timezone.now() - timedelta(days=periodo_dias)
    ocorrencias = OcorrenciaInspecao.objects.filter(criado_em__gte=desde)
    if bairro:
        ocorrencias = ocorrencias.filter(
            Q(local_ocorrencia__icontains=bairro)
            | Q(ambulante__licencas__ponto_ocupacao__bairro__iexact=bairro)
        ).distinct()
    ocorrencias = _aplicar_origem(ocorrencias, origem)
    if cidade or genero or faixa_etaria or escolaridade:
        ocorrencias = ocorrencias.filter(ambulante_id__in=ids_ambulantes or [0])

    retrato = _retrato_sociodemografico(ambulantes, hoje)
    cruzamentos = _cruzar_origem_local(ambulantes)

    bairros = list(
        PontoOcupacao.objects.filter(ativo=True)
        .exclude(bairro="")
        .order_by("bairro")
        .values_list("bairro", flat=True)
        .distinct()
    )
    cidades = list(
        Endereco.objects.exclude(cidade="")
        .order_by("cidade")
        .values_list("cidade", flat=True)
        .distinct()
    )

    payload = {
        "filtro_bairro": bairro,
        "filtro_origem": origem,
        "filtro_cidade": cidade,
        "filtro_genero": genero,
        "filtro_faixa": faixa_etaria,
        "filtro_escolaridade": escolaridade,
        "periodo_dias": periodo_dias,
        "bairros": bairros,
        "cidades": cidades,
        "generos": Genero.choices,
        "faixas_filtro": [(chave, rotulo) for chave, _min, _max, rotulo in FAIXAS_ETARIAS],
        "escolaridades_filtro": (
            ("fundamental", "Ensino Fundamental"),
            ("medio", "Ensino Médio"),
            ("superior", "Ensino Superior"),
        ),
        "licencas_ativas": ativas.count(),
        "licencas_vencidas": vencidas.count(),
        "licencas_pendentes": pendentes.count(),
        "licencas_suspensas": suspensas.count(),
        "licencas_vencendo": vencendo.count(),
        "ocorrencias_periodo": ocorrencias.count(),
        "residentes": ambulantes.filter(tipo_atuacao="fixo").count(),
        "itinerantes": ambulantes.filter(tipo_atuacao="movel").count(),
        "eventuais": ambulantes.filter(tipo_atuacao="eventual").count(),
        "ocupacao": ocupacao,
        "categorias": categorias,
        "alerta_lotacao": alerta_lotacao,
        "cidade_sede": CIDADE_SEDE,
        **retrato,
        **cruzamentos,
    }
    payload["graficos"] = dados_graficos(payload)
    return payload


def _filtrar_ambulantes(
    bairro="",
    origem="",
    cidade="",
    genero="",
    faixa_etaria="",
    escolaridade="",
):
    qs = Ambulante.objects.all()
    if bairro:
        qs = qs.filter(
            Q(licencas__ponto_ocupacao__bairro__iexact=bairro)
            | Q(ponto_pretendido__bairro__iexact=bairro)
        ).distinct()
    qs = _aplicar_origem(qs, origem, campo_ambulante="")
    if cidade:
        qs = qs.filter(enderecos__cidade__iexact=cidade).distinct()
    if genero:
        qs = qs.filter(genero=genero)
    if escolaridade == "fundamental":
        qs = qs.filter(
            Q(escolaridade__icontains="fundamental")
            | Q(escolaridade__in=("fundamental_incompleto", "fundamental_completo"))
        )
    elif escolaridade == "medio":
        qs = qs.filter(
            Q(escolaridade__icontains="medio")
            | Q(escolaridade__icontains="médio")
            | Q(escolaridade__in=("medio_incompleto", "medio_completo"))
        )
    elif escolaridade == "superior":
        qs = qs.filter(escolaridade__icontains="superior")
    if faixa_etaria:
        ids = [
            ambulante.pk
            for ambulante in qs.only("pk", "data_nasc")
            if _faixa_etaria(_idade(ambulante.data_nasc))[0] == faixa_etaria
        ]
        qs = qs.filter(pk__in=ids or [0])
    return qs


def _retrato_sociodemografico(ambulantes, hoje=None):
    hoje = hoje or timezone.localdate()
    registros = list(
        ambulantes.values(
            "pk",
            "data_nasc",
            "genero",
            "escolaridade",
            "nis",
            "renda_estimada",
            "num_funcionarios",
        )
    )
    total = len(registros)
    idades = Counter()
    generos = Counter()
    instrucao = Counter()
    rendas = Counter()
    com_nis = 0
    empregos = 0
    rendas_valores = []

    for item in registros:
        _chave, rotulo_idade = _faixa_etaria(_idade(item["data_nasc"], hoje))
        idades[rotulo_idade] += 1
        generos[item["genero"] or Genero.NAO_INFORMADO] += 1
        instrucao[_grupo_escolaridade(item["escolaridade"])] += 1
        if (item["nis"] or "").strip():
            com_nis += 1
        empregos += item["num_funcionarios"] or 0
        _chave_renda, rotulo_renda = _faixa_renda(item["renda_estimada"])
        rendas[rotulo_renda] += 1
        if item["renda_estimada"] is not None:
            rendas_valores.append(item["renda_estimada"])

    mapa_genero = dict(Genero.choices)
    dist_idade, _ = _contar(
        [{"nome": nome, "total": qtd} for nome, qtd in idades.items()]
    )
    dist_genero, _ = _contar(
        [
            {"nome": mapa_genero.get(chave, chave), "chave": chave, "total": qtd}
            for chave, qtd in generos.items()
        ]
    )
    dist_escolaridade, _ = _contar(
        [{"nome": nome, "total": qtd} for nome, qtd in instrucao.items()]
    )
    dist_renda, _ = _contar(
        [{"nome": nome, "total": qtd} for nome, qtd in rendas.items()]
    )
    media_renda = (
        round(sum(rendas_valores) / len(rendas_valores), 2) if rendas_valores else None
    )

    return {
        "total_ambulantes": total,
        "faixas_etarias": dist_idade,
        "distribuicao_genero": dist_genero,
        "distribuicao_escolaridade": dist_escolaridade,
        "faixas_renda": dist_renda,
        "com_nis": com_nis,
        "sem_nis": total - com_nis,
        "percentual_vulnerabilidade": _percentual(com_nis, total),
        "empregos_indiretos": empregos,
        "renda_media": media_renda,
        "com_renda_informada": len(rendas_valores),
    }


def _cruzar_origem_local(ambulantes):
    enderecos = {}
    for row in (
        Endereco.objects.filter(ambulante__in=ambulantes)
        .order_by("id")
        .values("ambulante_id", "cidade", "bairro")
    ):
        enderecos.setdefault(row["ambulante_id"], row)
    pontos = {}
    for row in (
        LicencaAlvara.objects.filter(
            ambulante__in=ambulantes, ponto_ocupacao__isnull=False
        )
        .order_by("-id")
        .values("ambulante_id", "ponto_ocupacao__bairro")
    ):
        pontos.setdefault(row["ambulante_id"], row["ponto_ocupacao__bairro"])
    pretendidos = {
        row["pk"]: row["ponto_pretendido__bairro"]
        for row in ambulantes.filter(ponto_pretendido__isnull=False).values(
            "pk", "ponto_pretendido__bairro"
        )
    }

    origem_cidade = Counter()
    origem_x_local = defaultdict(int)
    na_sede = 0
    de_fora = 0
    sem_endereco = 0

    for ambulante in ambulantes.only("pk"):
        endereco = enderecos.get(ambulante.pk)
        cidade = (endereco or {}).get("cidade") or "Não informado"
        bairro_atuacao = (
            pontos.get(ambulante.pk)
            or pretendidos.get(ambulante.pk)
            or "Não informado"
        )
        origem_cidade[cidade] += 1
        origem_x_local[(cidade, bairro_atuacao)] += 1
        if not endereco:
            sem_endereco += 1
        elif cidade.strip().lower() == CIDADE_SEDE.lower():
            na_sede += 1
        else:
            de_fora += 1

    cruzamento = [
        {
            "cidade_origem": cidade,
            "bairro_atuacao": bairro,
            "total": total,
        }
        for (cidade, bairro), total in sorted(
            origem_x_local.items(), key=lambda item: (-item[1], item[0][0], item[0][1])
        )
    ]
    cidades_origem, _ = _contar(
        [{"nome": nome, "total": qtd} for nome, qtd in origem_cidade.items()]
    )
    return {
        "cidades_origem": cidades_origem,
        "origem_x_local": cruzamento,
        "residentes_sede": na_sede,
        "residentes_fora": de_fora,
        "sem_endereco": sem_endereco,
    }


def indicadores_da_request(request):
    return indicadores_gerenciais(**filtros_do_request(request))


def indicadores_json(dados):
    """Recorte serializável para a API JSON (sem querysets)."""
    return {
        "filtros": {
            "bairro": dados["filtro_bairro"],
            "origem": dados["filtro_origem"] or "todos",
            "cidade": dados.get("filtro_cidade") or "todas",
            "genero": dados.get("filtro_genero") or "todos",
            "faixa_etaria": dados.get("filtro_faixa") or "todas",
            "escolaridade": dados.get("filtro_escolaridade") or "todas",
            "periodo_dias": dados["periodo_dias"],
        },
        "licencas": {
            "ativas": dados["licencas_ativas"],
            "vencidas": dados["licencas_vencidas"],
            "pendentes": dados["licencas_pendentes"],
            "suspensas": dados["licencas_suspensas"],
            "vencendo_30_dias": dados["licencas_vencendo"],
        },
        "ocorrencias_periodo": dados["ocorrencias_periodo"],
        "residentes": dados["residentes"],
        "itinerantes": dados["itinerantes"],
        "ocupacao": dados["ocupacao"],
        "categorias": dados["categorias"],
        "sociodemografico": {
            "total_ambulantes": dados["total_ambulantes"],
            "faixas_etarias": dados["faixas_etarias"],
            "genero": dados["distribuicao_genero"],
            "escolaridade": dados["distribuicao_escolaridade"],
            "renda": dados["faixas_renda"],
            "renda_media": (
                str(dados["renda_media"]) if dados["renda_media"] is not None else None
            ),
            "percentual_vulnerabilidade": dados["percentual_vulnerabilidade"],
            "empregos_indiretos": dados["empregos_indiretos"],
            "residentes_sede": dados["residentes_sede"],
            "residentes_fora": dados["residentes_fora"],
        },
        "origem_x_local": dados["origem_x_local"],
    }


def dados_graficos(dados):
    def pares(lista, chave="nome"):
        return {
            "labels": [item[chave] for item in lista],
            "data": [item["total"] for item in lista],
        }

    return {
        "categorias": pares(dados["categorias"]),
        "idades": pares(dados["faixas_etarias"]),
        "generos": pares(dados["distribuicao_genero"]),
        "escolaridade": pares(dados["distribuicao_escolaridade"]),
        "renda": pares(dados["faixas_renda"]),
        "origem": {
            "labels": ["Vitória da Conquista", "Outros municípios", "Sem endereço"],
            "data": [
                dados["residentes_sede"],
                dados["residentes_fora"],
                dados["sem_endereco"],
            ],
        },
    }


def linhas_ambulantes(dados_filtro):
    qs = (
        _filtrar_ambulantes(
            bairro=dados_filtro.get("filtro_bairro") or dados_filtro.get("bairro", ""),
            origem=dados_filtro.get("filtro_origem") or dados_filtro.get("origem", ""),
            cidade=dados_filtro.get("filtro_cidade") or dados_filtro.get("cidade", ""),
            genero=dados_filtro.get("filtro_genero") or dados_filtro.get("genero", ""),
            faixa_etaria=dados_filtro.get("filtro_faixa")
            or dados_filtro.get("faixa_etaria", ""),
            escolaridade=dados_filtro.get("filtro_escolaridade")
            or dados_filtro.get("escolaridade", ""),
        )
        .prefetch_related("enderecos", "licencas__ponto_ocupacao")
        .select_related("ponto_pretendido")
        .order_by("nome", "sobrenome")
    )
    hoje = timezone.localdate()
    linhas = []
    for ambulante in qs:
        endereco = ambulante.enderecos.order_by("id").first()
        licenca = ambulante.licencas.order_by("-id").first()
        ponto = (
            licenca.ponto_ocupacao
            if licenca and licenca.ponto_ocupacao_id
            else ambulante.ponto_pretendido
        )
        linhas.append(
            {
                "nome": ambulante.nome_completo,
                "cpf": ambulante.cpf_formatado,
                "genero": ambulante.get_genero_display(),
                "idade": _idade(ambulante.data_nasc, hoje),
                "escolaridade": _grupo_escolaridade(ambulante.escolaridade),
                "nis": "Sim" if (ambulante.nis or "").strip() else "Não",
                "renda": ambulante.renda_estimada,
                "auxiliares": ambulante.num_funcionarios,
                "tipo_atuacao": ambulante.tipo_atuacao or "",
                "cidade_origem": endereco.cidade if endereco else "",
                "bairro_origem": endereco.bairro if endereco else "",
                "bairro_atuacao": ponto.bairro if ponto else "",
                "score": ambulante.pontuacao,
            }
        )
    return linhas


def montar_planilha(dados):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    cabecalho = Font(bold=True, color="FFFFFF")
    fundo = PatternFill("solid", fgColor="2563EB")

    def escrever(planilha, titulo, cabecalhos, linhas):
        planilha.append([titulo])
        planilha.merge_cells(
            start_row=1, start_column=1, end_row=1, end_column=max(len(cabecalhos), 1)
        )
        planilha.append(cabecalhos)
        for coluna, _nome in enumerate(cabecalhos, start=1):
            celula = planilha.cell(row=2, column=coluna)
            celula.font = cabecalho
            celula.fill = fundo
            celula.alignment = Alignment(horizontal="center")
        for linha in linhas:
            planilha.append(list(linha))
        for indice, _nome in enumerate(cabecalhos, start=1):
            planilha.column_dimensions[get_column_letter(indice)].width = 22

    resumo = workbook.active
    resumo.title = "Resumo"
    escrever(
        resumo,
        "Indicadores APAL — Vitória da Conquista",
        ["Indicador", "Valor"],
        [
            ("Licenças ativas", dados["licencas_ativas"]),
            ("Licenças vencidas", dados["licencas_vencidas"]),
            ("Pendentes na fila", dados["licencas_pendentes"]),
            ("Infrações no período", dados["ocorrencias_periodo"]),
            ("Ambulantes no recorte", dados["total_ambulantes"]),
            ("Com NIS / CadÚnico", f"{dados['percentual_vulnerabilidade']}%"),
            ("Empregos indiretos", dados["empregos_indiretos"]),
            (
                "Renda média (R$)",
                dados["renda_media"] if dados["renda_media"] is not None else "—",
            ),
            ("Moram na sede", dados["residentes_sede"]),
            ("Moram em outro município", dados["residentes_fora"]),
        ],
    )

    ocupacao = workbook.create_sheet("Ocupação")
    escrever(
        ocupacao,
        "Ocupação de vagas por bairro",
        ["Bairro", "Ocupadas", "Total", "Lotação %", "Alvarás ativos"],
        [
            (
                item["bairro"],
                item["ocupados"],
                item["total"],
                item["percentual"],
                item["licencas_ativas"],
            )
            for item in dados["ocupacao"]
        ],
    )

    categorias = workbook.create_sheet("Categorias")
    escrever(
        categorias,
        "Licenças ativas por categoria",
        ["Categoria", "Total"],
        [(item["nome"], item["total"]) for item in dados["categorias"]],
    )

    socio = workbook.create_sheet("Sociodemográfico")
    socio.append(["Faixa etária", "Total", "%"])
    for item in dados["faixas_etarias"]:
        socio.append([item["nome"], item["total"], item["percentual"]])
    socio.append([])
    socio.append(["Gênero", "Total", "%"])
    for item in dados["distribuicao_genero"]:
        socio.append([item["nome"], item["total"], item["percentual"]])
    socio.append([])
    socio.append(["Escolaridade", "Total", "%"])
    for item in dados["distribuicao_escolaridade"]:
        socio.append([item["nome"], item["total"], item["percentual"]])
    socio.append([])
    socio.append(["Faixa de renda", "Total", "%"])
    for item in dados["faixas_renda"]:
        socio.append([item["nome"], item["total"], item["percentual"]])

    cruzado = workbook.create_sheet("Origem x local")
    escrever(
        cruzado,
        "Cidade de moradia × bairro de atuação",
        ["Cidade de origem", "Bairro de atuação", "Ambulantes"],
        [
            (item["cidade_origem"], item["bairro_atuacao"], item["total"])
            for item in dados["origem_x_local"]
        ],
    )

    detalhe = workbook.create_sheet("Ambulantes")
    escrever(
        detalhe,
        "Base nominativa do recorte",
        [
            "Nome",
            "CPF",
            "Gênero",
            "Idade",
            "Escolaridade",
            "NIS",
            "Renda (R$)",
            "Auxiliares",
            "Atuação",
            "Cidade origem",
            "Bairro origem",
            "Bairro atuação",
            "Score",
        ],
        [
            (
                linha["nome"],
                linha["cpf"],
                linha["genero"],
                linha["idade"] if linha["idade"] is not None else "",
                linha["escolaridade"],
                linha["nis"],
                linha["renda"] if linha["renda"] is not None else "",
                linha["auxiliares"],
                linha["tipo_atuacao"],
                linha["cidade_origem"],
                linha["bairro_origem"],
                linha["bairro_atuacao"],
                linha["score"],
            )
            for linha in linhas_ambulantes(dados)
        ],
    )

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer
